from __future__ import annotations

import io
import json
import os
import re
from typing import Any, Literal

import httpx
import matplotlib
import sympy as sp
from fastapi import Depends, FastAPI, File, Header, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import Response
from pydantic import BaseModel, Field

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402

app = FastAPI(title="LearnFlow 工具服务", version="2.0.0")
app.add_middleware(
    CORSMiddleware,
    allow_origins=[origin.strip() for origin in os.getenv("ALLOWED_ORIGINS", "http://localhost:3000").split(",")],
    allow_credentials=False,
    allow_methods=["POST", "GET"],
    allow_headers=["Authorization", "Content-Type"],
)


def require_token(authorization: str | None = Header(default=None)) -> None:
    expected = os.getenv("TOOL_SERVICE_TOKEN", "")
    if expected and authorization != f"Bearer {expected}":
        raise HTTPException(status_code=401, detail="工具服务令牌无效")


SYMBOLS = {name: sp.Symbol(name) for name in ["x", "y", "z", "t", "n", "k", "a", "b", "c"]}
SAFE_GLOBALS: dict[str, Any] = {
    "__builtins__": {},
    "Symbol": sp.Symbol,
    "Integer": sp.Integer,
    "Float": sp.Float,
    "Rational": sp.Rational,
    "sin": sp.sin,
    "cos": sp.cos,
    "tan": sp.tan,
    "asin": sp.asin,
    "acos": sp.acos,
    "atan": sp.atan,
    "exp": sp.exp,
    "log": sp.log,
    "sqrt": sp.sqrt,
    "Abs": sp.Abs,
    "factorial": sp.factorial,
    "pi": sp.pi,
    "E": sp.E,
    "oo": sp.oo,
    **SYMBOLS,
}


def parse_math(value: str) -> sp.Expr:
    if not value.strip() or len(value) > 3000:
        raise HTTPException(status_code=400, detail="数学表达式为空或过长")
    if re.search(r"__|\b(import|exec|eval|open|compile|lambda|globals|locals)\b|[;\[\]{}]", value):
        raise HTTPException(status_code=400, detail="表达式包含不允许的内容")
    try:
        return sp.sympify(value, locals=SAFE_GLOBALS, evaluate=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"无法解析数学表达式：{exc}") from exc


class MathRequest(BaseModel):
    action: Literal["simplify", "derivative", "integral", "limit", "solve", "verify", "substitute", "units"]
    expression: str
    second_expression: str | None = None
    variable: str = "x"
    point: str | None = None
    values: dict[str, float] = Field(default_factory=dict)
    units_to: str | None = None


@app.get("/health")
def health() -> dict[str, str]:
    return {"status": "ok", "service": "zhijie-tools", "version": "2.0.0"}


@app.post("/math/evaluate", dependencies=[Depends(require_token)])
def math_evaluate(request: MathRequest) -> dict[str, Any]:
    variable = SYMBOLS.get(request.variable, sp.Symbol(request.variable))
    expression = parse_math(request.expression)
    try:
        if request.action == "simplify":
            result = sp.simplify(expression)
        elif request.action == "derivative":
            result = sp.diff(expression, variable)
        elif request.action == "integral":
            result = sp.integrate(expression, variable)
        elif request.action == "limit":
            if request.point is None:
                raise HTTPException(status_code=400, detail="求极限需要 point")
            result = sp.limit(expression, variable, parse_math(request.point))
        elif request.action == "solve":
            right = parse_math(request.second_expression or "0")
            result = sp.solve(sp.Eq(expression, right), variable)
        elif request.action == "verify":
            if request.second_expression is None:
                raise HTTPException(status_code=400, detail="等价验证需要 second_expression")
            right = parse_math(request.second_expression)
            difference = sp.simplify(expression - right)
            return {
                "verified": difference == 0,
                "difference": sp.sstr(difference),
                "latex": sp.latex(difference),
                "method": "SymPy symbolic equivalence",
            }
        elif request.action == "substitute":
            substitutions = {SYMBOLS.get(name, sp.Symbol(name)): value for name, value in request.values.items()}
            result = sp.N(expression.subs(substitutions))
        elif request.action == "units":
            try:
                import pint
            except ImportError as exc:
                raise HTTPException(status_code=503, detail="工具服务尚未安装 pint") from exc
            registry = pint.UnitRegistry()
            quantity = registry(request.expression)
            result = quantity.to(request.units_to) if request.units_to else quantity.to_base_units()
            return {"result": str(result), "method": "Pint unit conversion"}
        else:
            raise HTTPException(status_code=400, detail="未知数学动作")
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"符号计算失败：{exc}") from exc
    return {
        "result": sp.sstr(result),
        "latex": sp.latex(result),
        "method": "SymPy",
        "action": request.action,
    }


class PlotRequest(BaseModel):
    expression: str
    variable: str = "x"
    minimum: float = -10
    maximum: float = 10
    samples: int = Field(default=800, ge=100, le=5000)
    title: str | None = None


@app.post("/math/plot", dependencies=[Depends(require_token)])
def math_plot(request: PlotRequest) -> Response:
    import numpy as np

    variable = SYMBOLS.get(request.variable, sp.Symbol(request.variable))
    expression = parse_math(request.expression)
    function = sp.lambdify(variable, expression, modules=["numpy"])
    xs = np.linspace(request.minimum, request.maximum, request.samples)
    try:
        ys = function(xs)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"函数无法在指定区间绘图：{exc}") from exc
    figure, axis = plt.subplots(figsize=(9, 5), dpi=160)
    axis.plot(xs, ys, color="#6547f5", linewidth=2)
    axis.axhline(0, color="#9ca3af", linewidth=0.8)
    axis.axvline(0, color="#9ca3af", linewidth=0.8)
    axis.grid(alpha=0.2)
    axis.set_title(request.title or f"{request.expression}")
    axis.set_xlabel(request.variable)
    figure.tight_layout()
    output = io.BytesIO()
    figure.savefig(output, format="png", bbox_inches="tight")
    plt.close(figure)
    return Response(output.getvalue(), media_type="image/png", headers={"X-Tool": "SymPy+Matplotlib"})


def parse_pdf(data: bytes) -> tuple[str, int]:
    import fitz

    document = fitz.open(stream=data, filetype="pdf")
    pages = []
    for index, page in enumerate(document):
        pages.append(f"\n\n<!-- Page {index + 1} -->\n\n{page.get_text('text')}")
    return "".join(pages), len(document)


def parse_presentation(data: bytes) -> tuple[str, int]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(data))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides):
        content = [f"\n\n<!-- Page {index + 1} -->\n\n# 幻灯片 {index + 1}"]
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                content.append(shape.text.strip())
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    content.append(" | ".join(cell.text.strip() for cell in row.cells))
        slides.append("\n\n".join(content))
    return "\n".join(slides), len(presentation.slides)


def parse_word(data: bytes) -> tuple[str, int | None]:
    from docx import Document

    document = Document(io.BytesIO(data))
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n\n".join(parts), None


@app.post("/parse", dependencies=[Depends(require_token)])
async def parse_document(file: UploadFile = File(...)) -> dict[str, Any]:
    data = await file.read()
    if len(data) > 60 * 1024 * 1024:
        raise HTTPException(status_code=413, detail="工具服务限制单文件 60 MB")
    name = (file.filename or "material").lower()
    try:
        if name.endswith(".pdf"):
            text, page_count = parse_pdf(data)
        elif name.endswith((".ppt", ".pptx")):
            text, page_count = parse_presentation(data)
        elif name.endswith((".doc", ".docx")):
            text, page_count = parse_word(data)
        elif name.endswith((".txt", ".md", ".markdown")):
            text, page_count = data.decode("utf-8", errors="replace"), None
        else:
            raise HTTPException(status_code=415, detail="本地结构解析支持 PDF、PPTX、DOCX、TXT 和 Markdown；图片请使用主模型 OCR。")
    except HTTPException:
        raise
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=422, detail=f"文档结构解析失败：{exc}") from exc
    if not text.strip():
        raise HTTPException(status_code=422, detail="文档中没有提取出可用文字；扫描件请使用主模型 OCR。")
    return {"markdown": text, "pageCount": page_count, "parser": "local-structure-parser"}


class JudgeRequest(BaseModel):
    source_code: str
    language_id: int
    stdin: str = ""
    expected_output: str | None = None
    cpu_time_limit: float = 4
    memory_limit: int = 256000


@app.post("/oj/run", dependencies=[Depends(require_token)])
async def run_oj(request: JudgeRequest) -> dict[str, Any]:
    base_url = os.getenv("JUDGE0_BASE_URL", "").rstrip("/")
    if not base_url:
        raise HTTPException(status_code=503, detail="尚未配置 JUDGE0_BASE_URL，代码不会在工具服务进程中直接执行。")
    headers = {"Content-Type": "application/json"}
    if os.getenv("JUDGE0_API_KEY"):
        headers["X-Auth-Token"] = os.getenv("JUDGE0_API_KEY", "")
    payload = request.model_dump()
    async with httpx.AsyncClient(timeout=45) as client:
        created = await client.post(f"{base_url}/submissions?base64_encoded=false&wait=true", headers=headers, content=json.dumps(payload))
    if created.status_code >= 400:
        raise HTTPException(status_code=502, detail=f"Judge0 调用失败：{created.text[:300]}")
    result = created.json()
    return {
        "stdout": result.get("stdout"),
        "stderr": result.get("stderr"),
        "compile_output": result.get("compile_output"),
        "message": result.get("message"),
        "time": result.get("time"),
        "memory": result.get("memory"),
        "status": result.get("status"),
        "source": "Judge0",
    }
